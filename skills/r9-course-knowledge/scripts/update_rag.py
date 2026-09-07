#!/usr/bin/env python3
"""
更新 R9 课程知识库的 references 文本。

用法：
    python3 scripts/update_rag.py

脚本会读取 /Users/r9/Desktop/核心课程大课（材料）/ 下的课件（PDF/PPTX），
提取文本后写入 references/ 目录下对应的主题 markdown 文件。
"""

import fitz
import os
import re
from pptx import Presentation

BASE_DIR = "/Users/r9/Desktop/核心课程大课（材料）"
OUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "references")

THEME_FOLDERS = {
    "01-财富管理业务与行业演变": "01-wealth-management-business.md",
    "02-顾问式服务与组合管理": "02-advisory-service.md",
    "03-基金销售方法论": "03-fund-sales-methodology.md",
    "04-资产配置理论与实践": "04-asset-allocation.md",
    "05-基金评价与研究": "05-fund-research.md",
    "06-客户陪伴与售后服务": "06-client-companion.md",
    "07-市场解读与晨夕会": "07-market-briefing.md",
    "08-AI与科技赋能": "08-ai-empowerment.md",
    "09-投资者教育与行为": "09-investor-education.md",
}


def clean_text(text: str) -> str:
    text = text.strip()
    text = re.sub(r"\n+", "\n", text)
    return text


def extract_pdf(path: str) -> str:
    try:
        doc = fitz.open(path)
    except Exception as e:
        return f"**无法读取 PDF：{e}**\n"
    parts = []
    for i, page in enumerate(doc, 1):
        text = page.get_text().strip()
        if text:
            parts.append(f"### 第{i}页\n\n{clean_text(text)}\n")
    return "\n".join(parts)


def extract_pptx(path: str, max_slides: int = None) -> str:
    try:
        prs = Presentation(path)
    except Exception as e:
        return f"**无法读取 PPT：{e}**\n"
    slides = list(prs.slides)
    if max_slides:
        slides = slides[:max_slides]
    parts = []
    for i, slide in enumerate(slides, 1):
        texts = []
        for shape in slide.shapes:
            try:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
            except Exception:
                continue
        slide_text = "\n".join(texts)
        if slide_text.strip():
            parts.append(f"### 第{i}页\n\n{clean_text(slide_text)}\n")
    return "\n".join(parts)


def update_theme(folder: str, ref_name: str):
    folder_path = os.path.join(BASE_DIR, folder)
    if not os.path.isdir(folder_path):
        print(f"跳过：文件夹不存在 {folder_path}")
        return

    ref_path = os.path.join(OUT_DIR, ref_name)
    with open(ref_path, "w", encoding="utf-8") as out_f:
        out_f.write(f"# {folder}\n\n")

        files = sorted(os.listdir(folder_path))
        for fname in files:
            if fname.startswith("~$") or fname == ".DS_Store":
                continue
            fpath = os.path.join(folder_path, fname)
            if not os.path.isfile(fpath):
                continue
            _, ext = os.path.splitext(fname)
            ext = ext.lower()
            if ext not in (".pdf", ".pptx"):
                continue

            print(f"  处理：{folder}/{fname}")
            out_f.write(f"## 课件：{fname}\n\n")
            if ext == ".pdf":
                out_f.write(extract_pdf(fpath))
            elif ext == ".pptx":
                max_s = 300 if "基金配置能力建设" in fname else None
                out_f.write(extract_pptx(fpath, max_slides=max_s))
            out_f.write("\n---\n\n")
    print(f"已更新：{ref_path}")


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    for folder, ref_name in THEME_FOLDERS.items():
        print(f"\n更新主题：{folder}")
        update_theme(folder, ref_name)
    print("\n全部更新完成。")


if __name__ == "__main__":
    main()
